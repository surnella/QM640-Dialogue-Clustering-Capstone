"""
APA 7 (student paper) document helpers built on python-docx.

Uniform rules applied everywhere:
  Times New Roman 12 pt, double spacing, 0 pt before/after, 1-inch margins,
  flush-left / ragged-right text, 0.5-inch first-line indents, page number top right,
  Level 1 = centered bold, Level 2 = flush-left bold, Level 3 = flush-left bold italic,
  tables with horizontal rules only (top, below header, bottom), one table-body size (10 pt).
Inline markup supported in text: <i>..</i>, <b>..</b>, <sub>..</sub>, <sup>..</sup>.
"""

import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
BODY_PT = 12
TABLE_PT = 10
_TAG = re.compile(r"(</?(?:i|b|sub|sup)>)")
_URL = re.compile(r"(https?://[^\s<>]+[^\s<>.,;:)\]])")
# In-text table/figure citations: "Table 3", "Tables 6 and 7", "Tables 3, 4, and 5", "Figure E5", "Tables D1–D3"
_NUM = r"[A-E]?\d+"
_XREF = re.compile(rf"\b(Tables?|Figures?)\s+({_NUM}(?:(?:,\s+and\s+|,\s+|\s+and\s+|–){_NUM})*)(?!\w)")
_BM_ID = [0]


def _bm_names(kind, label):
    """Hidden bookmark names (Word convention: leading underscore) for a caption label."""
    k = "Tab" if kind.startswith("Table") else "Fig"
    return f"_Ref{k}{label}", f"_Ref{k}N{label}"  # whole label ("Table 3"), number only ("3")


def _bookmark(par, name, start=True):
    """Open or close a bookmark at the current end of the paragraph."""
    el = OxmlElement("w:bookmarkStart" if start else "w:bookmarkEnd")
    if start:
        _BM_ID[0] += 1
        _BM_OPEN[name] = _BM_ID[0]
        el.set(qn("w:name"), name)
    el.set(qn("w:id"), str(_BM_ID[0] if start else _BM_OPEN.pop(name)))
    par._p.append(el)


def _caption_label(par, kind, label):
    """Bold caption label ("Table 3") wrapped in the bookmarks that cross-references point to."""
    whole, num = _bm_names(kind, label)
    _bookmark(par, whole)
    _set_font(par.add_run(f"{kind} "), bold=True)
    _bookmark(par, num)
    _set_font(par.add_run(str(label)), bold=True)
    _bookmark(par, num, start=False)
    _bookmark(par, whole, start=False)
    _registry(par)["targets"].update((whole, num))


def _registry(par):
    part = par.part
    if not hasattr(part, "_xref_registry"):
        part._xref_registry = {"targets": set(), "refs": {}}
    return part._xref_registry


_BM_OPEN = {}


def _ref_field(par, bookmark, text, size, bold, italic, sub, sup):
    """Insert a Word cross-reference (REF field with \\h = clickable) whose cached result is `text`."""
    def fld(kind):
        r = OxmlElement("w:r")
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), kind)
        r.append(f)
        par._p.append(r)
    fld("begin")
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" REF {bookmark} \\h \\* MERGEFORMAT "
    r.append(it)
    par._p.append(r)
    fld("separate")
    run = par.add_run(text)
    _set_font(run, size, bold=bold, italic=italic)
    if sub:
        run.font.subscript = True
    if sup:
        run.font.superscript = True
    fld("end")
    refs = _registry(par)["refs"]
    refs[bookmark] = refs.get(bookmark, 0) + 1


def _plain_run(par, text, size, st):
    r = par.add_run(text)
    _set_font(r, size, bold=st["b"], italic=st["i"])
    if st["sub"]:
        r.font.subscript = True
    if st["sup"]:
        r.font.superscript = True


def _xref_runs(par, text, size, st):
    """Write text, turning every table/figure citation into a clickable cross-reference."""
    pos = 0
    args = (size, st["b"], st["i"], st["sub"], st["sup"])
    for m in _XREF.finditer(text):
        if m.start() > pos:
            _plain_run(par, text[pos:m.start()], size, st)
        word, nums = m.group(1), m.group(2)
        labels = re.findall(_NUM, nums)
        if len(labels) == 1 and not word.endswith("s"):
            whole, _ = _bm_names(word, labels[0])
            _ref_field(par, whole, m.group(0), *args)
        else:
            _plain_run(par, text[m.start():m.start(2)], size, st)
            p = 0
            for nm in re.finditer(_NUM, nums):
                if nm.start() > p:
                    _plain_run(par, nums[p:nm.start()], size, st)
                _, num_bm = _bm_names(word, nm.group(0))
                _ref_field(par, num_bm, nm.group(0), *args)
                p = nm.end()
        pos = m.end()
    if pos < len(text):
        _plain_run(par, text[pos:], size, st)


def _hyperlink_run(par, url, text, size, bold, italic, sub, sup):
    """Append an external hyperlink (clickable in Word and in the exported PDF)."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    r_id = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    hl.set(qn("w:history"), "1")
    r = par.add_run(text)
    _set_font(r, size, bold=bold, italic=italic)
    if sub:
        r.font.subscript = True
    if sup:
        r.font.superscript = True
    r._r.getparent().remove(r._r)
    hl.append(r._r)
    par._p.append(hl)


def _set_font(run, size=BODY_PT, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_runs(par, text, size=BODY_PT, bold=False, italic=False, xref=True):
    """Add text with inline <i>/<b>/<sub>/<sup> markup to a paragraph.

    URLs become external hyperlinks; with xref=True, "Table N"/"Figure N" citations become clickable
    cross-references to the caption bookmarks.
    """
    state = dict(i=italic, b=bold, sub=False, sup=False)
    for tok in _TAG.split(text):
        if not tok:
            continue
        m = re.fullmatch(r"<(/?)(i|b|sub|sup)>", tok)
        if m:
            state[m.group(2)] = (m.group(1) == "")
            continue
        for j, piece in enumerate(_URL.split(tok)):
            if not piece:
                continue
            if j % 2 == 1:
                _hyperlink_run(par, piece, piece, size, state["b"], state["i"], state["sub"], state["sup"])
                continue
            if xref:
                _xref_runs(par, piece, size, state)
            else:
                _plain_run(par, piece, size, state)
    return par


def _fmt(par, align=WD_ALIGN_PARAGRAPH.LEFT, first=0.0, left=0.0, spacing=2.0, before=0, after=0,
         keep_next=False):
    pf = par.paragraph_format
    par.alignment = align
    pf.first_line_indent = Inches(first)
    pf.left_indent = Inches(left)
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = keep_next
    pf.widow_control = True
    return par


class APADoc:
    def __init__(self):
        self.doc = Document()
        self.fig_n = 0
        self.tab_n = 0
        self.headings = []  # (level, text) in order, for the table of contents
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(sec, side, Inches(1))
        sec.header_distance = Inches(0.5)
        sec.footer_distance = Inches(0.5)
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(BODY_PT)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.line_spacing = 2.0
        st.paragraph_format.space_before = Pt(0)
        st.paragraph_format.space_after = Pt(0)
        for lvl in (1, 2, 3):
            hs = self.doc.styles[f"Heading {lvl}"]
            hs.font.name, hs.font.size = FONT, Pt(BODY_PT)
            hs.font.color.rgb = RGBColor(0, 0, 0)
            hs.font.bold = True
            hs.font.italic = (lvl == 3)
            rpr = hs.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts"); rpr.append(rf)
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia", "w:asciiTheme", "w:hAnsiTheme"):
                if a.endswith("Theme"):
                    if rf.get(qn(a)) is not None:
                        del rf.attrib[qn(a)]
                else:
                    rf.set(qn(a), FONT)
            pf = hs.paragraph_format
            pf.line_spacing, pf.space_before, pf.space_after = 2.0, Pt(0), Pt(0)
            pf.keep_with_next = True
        self._page_number_header(sec)

    # ------------------------------------------------------------------ layout
    def _page_number_header(self, sec):
        p = sec.header.paragraphs[0]
        _fmt(p, align=WD_ALIGN_PARAGRAPH.RIGHT, spacing=1.0)
        run = p.add_run()
        _set_font(run)
        for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
            else:
                el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = txt
            run._r.append(el)

    def page_break(self):
        p = self.doc.add_paragraph()
        _fmt(p, spacing=1.0)
        p.add_run().add_break(WD_BREAK.PAGE)

    def blank(self):
        _fmt(self.doc.add_paragraph())

    # ------------------------------------------------------------------ text
    def title_page(self, title_lines, info_lines):
        for _ in range(3):
            self.blank()
        for t in title_lines:
            add_runs(_fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER), t, bold=True, xref=False)
        self.blank()
        for t in info_lines:
            add_runs(_fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER), t, xref=False)
        self.page_break()

    def heading(self, text, level=1, toc=True, new_page=False):
        p = self.doc.add_paragraph(style=f"Heading {level}")
        _fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER if level == 1 else WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
        p.paragraph_format.page_break_before = new_page
        add_runs(p, text, bold=True, italic=(level == 3), xref=False)
        if toc:
            self.headings.append((level, re.sub(r"</?\w+>", "", text)))
        return p

    def para(self, text, indent=True, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=False):
        p = _fmt(self.doc.add_paragraph(), align=align, first=0.5 if indent else 0.0, keep_next=keep_next)
        return add_runs(p, text)

    def lead_para(self, lead, text):
        """Paragraph opening with an italic run-in label, e.g., an assumption name."""
        return self.para(f"<i>{lead}.</i> {text}")

    def items(self, entries, numbered=True):
        """APA seriation as indented, double-spaced, hanging-indent paragraphs."""
        for i, t in enumerate(entries, 1):
            label = f"{i}." if numbered else "\u2022"
            p = _fmt(self.doc.add_paragraph(), first=-0.25, left=0.75)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(0.75))
            add_runs(p, f"{label}\t{t}")

    def equation(self, text, label=None):
        p = _fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
        add_runs(p, text + (f"    ({label})" if label else ""))

    def reference(self, text):
        p = _fmt(self.doc.add_paragraph(), first=-0.5, left=0.5)
        add_runs(p, text, xref=False)

    # ------------------------------------------------------------------ TOC
    def toc(self, entries=None, pages=None):
        p = _fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER)
        add_runs(p, "Contents", bold=True, xref=False)
        for idx, (lvl, txt) in enumerate(entries or []):
            if lvl > 2:
                continue
            q = _fmt(self.doc.add_paragraph(), left=0.5 * (lvl - 1))
            q.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            pg = (pages or {}).get(idx, "")
            add_runs(q, f"{txt}\t{pg}", xref=False)
        self.page_break()

    # ------------------------------------------------------------------ tables & figures
    def _borders(self, table):
        tblPr = table._tbl.tblPr
        b = OxmlElement("w:tblBorders")
        for edge, val in (("top", "single"), ("bottom", "single"), ("left", "nil"), ("right", "nil"),
                          ("insideH", "nil"), ("insideV", "nil")):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), val)
            if val != "nil":
                e.set(qn("w:sz"), "8"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
            b.append(e)
        tblPr.append(b)

    @staticmethod
    def _cell_border_bottom(cell):
        tcPr = cell._tc.get_or_add_tcPr()
        bd = OxmlElement("w:tcBorders")
        e = OxmlElement("w:bottom")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "6"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "000000")
        bd.append(e)
        tcPr.append(bd)

    def table(self, title, header, rows, note=None, widths=None, align=None, label=None, split=False):
        """APA table: bold number, italic title, rules only, consistent 10-pt body, note below."""
        if label is None:
            self.tab_n += 1
            label = str(self.tab_n)
        _caption_label(_fmt(self.doc.add_paragraph(), before=12, keep_next=True), "Table", label)
        add_runs(_fmt(self.doc.add_paragraph(), keep_next=True), title, italic=True, xref=False)
        ncol = len(header)
        widths = widths or [6.5 / ncol] * ncol
        align = align or (["l"] + ["c"] * (ncol - 1))
        t = self.doc.add_table(rows=1 + len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._borders(t)
        amap = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
        n_rows = 1 + len(rows)
        keep_all = n_rows <= 18 and not split  # short tables stay on one page; long ones keep header + 2 rows
        for r_i, row in enumerate([header] + rows):
            tr = t.rows[r_i]
            trPr = tr._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trPr.append(cs)
            if r_i == 0:
                th = OxmlElement("w:tblHeader"); trPr.append(th)
            keep = ((r_i < n_rows - 1) and (keep_all or r_i < 2)) or (r_i == n_rows - 1 and bool(note))
            for c_i, val in enumerate(row):
                cell = tr.cells[c_i]
                cell.width = Inches(widths[c_i])
                p = cell.paragraphs[0]
                _fmt(p, align=WD_ALIGN_PARAGRAPH.CENTER if r_i == 0 else amap[align[c_i]], spacing=1.0,
                     before=2, after=2, keep_next=keep)
                add_runs(p, str(val), size=TABLE_PT)
                if r_i == 0:
                    self._cell_border_bottom(cell)
        if note:
            add_runs(_fmt(self.doc.add_paragraph(), after=12), f"<i>Note.</i> {note}")
        else:
            _fmt(self.doc.add_paragraph(), spacing=1.0, after=6)
        return label

    def figure(self, title, image_path, note=None, width=6.0, label=None):
        if label is None:
            self.fig_n += 1
            label = str(self.fig_n)
        _caption_label(_fmt(self.doc.add_paragraph(), before=12, keep_next=True), "Figure", label)
        add_runs(_fmt(self.doc.add_paragraph(), keep_next=True), title, italic=True, xref=False)
        p = _fmt(self.doc.add_paragraph(), align=WD_ALIGN_PARAGRAPH.CENTER, spacing=1.0, keep_next=True)
        p.add_run().add_picture(image_path, width=Inches(width))
        if note:
            add_runs(_fmt(self.doc.add_paragraph(), after=12), f"<i>Note.</i> {note}")
        return label

    def xref_summary(self):
        """(number of cross-reference fields, number of caption targets, references without a target)."""
        reg = getattr(self.doc.part, "_xref_registry", {"targets": set(), "refs": {}})
        missing = sorted(set(reg["refs"]) - reg["targets"])
        return sum(reg["refs"].values()), len(reg["targets"]) // 2, missing

    def save(self, path):
        n_refs, n_targets, missing = self.xref_summary()
        if missing:
            raise ValueError(f"Cross-references without a matching caption: {missing}")
        self.doc.save(path)
